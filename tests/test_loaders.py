import hashlib
import json
import logging
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

from hybrid_rag.ingestion.loaders import (
    DocumentLoadError,
    discover_documents,
    load_document,
    load_documents,
)


def make_text_pdf(path: Path, pages: tuple[str, ...]) -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
        }
    )
    font_reference = writer._add_object(font)
    for text in pages:
        page = writer.add_blank_page(width=300, height=300)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as destination:
        writer.write(destination)


def make_image_only_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    image = DecodedStreamObject()
    image.set_data(b"\x00\x00\x00")
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1),
            NameObject("/Height"): NumberObject(1),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    image_reference = writer._add_object(image)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/XObject"): DictionaryObject(
                {NameObject("/Im0"): image_reference}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"q 100 0 0 100 0 0 cm /Im0 Do Q")
    page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as destination:
        writer.write(destination)


def make_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Paragraph before table")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Value A"
    table.cell(1, 1).text = "Value B"
    document.add_paragraph("Paragraph after table")
    document.save(path)


def test_pdf_loading_preserves_pages_hash_and_relative_source(tmp_path: Path) -> None:
    documents_dir = tmp_path / "documents"
    nested = documents_dir / "manuals"
    nested.mkdir(parents=True)
    source = nested / "guide.pdf"
    make_text_pdf(source, ("First PDF page.", "Second PDF page."))

    record = load_document(source, documents_dir)

    assert record.source_path == "manuals/guide.pdf"
    assert record.file_name == "guide.pdf"
    assert record.file_type == "pdf"
    assert record.file_hash == hashlib.sha256(source.read_bytes()).hexdigest()
    assert [(segment.page_number, segment.text.strip()) for segment in record.segments] == [
        (1, "First PDF page."),
        (2, "Second PDF page."),
    ]
    assert load_document(Path("manuals/guide.pdf"), documents_dir).document_id == record.document_id


def test_docx_loading_preserves_paragraph_and_table_order(tmp_path: Path) -> None:
    documents_dir = tmp_path / "documents"
    documents_dir.mkdir()
    source = documents_dir / "table.docx"
    make_docx(source)

    record = load_document(source, documents_dir)

    assert record.file_type == "docx"
    assert [segment.text for segment in record.segments] == [
        "Paragraph before table",
        "Header A | Header B",
        "Value A | Value B",
        "Paragraph after table",
    ]


def test_discovery_is_recursive_and_skips_temporary_and_unsupported_files(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    documents_dir = tmp_path / "documents"
    nested = documents_dir / "nested"
    nested.mkdir(parents=True)
    (nested / "valid.PDF").write_bytes(b"pdf")
    (documents_dir / "notes.txt").write_text("not a source document", encoding="utf-8")
    (documents_dir / "~$draft.docx").write_bytes(b"temporary")
    (documents_dir / "partial.pdf.part").write_bytes(b"partial")
    (documents_dir / ".gitkeep").write_text("", encoding="utf-8")

    with caplog.at_level(logging.INFO, logger="hybrid_rag.ingestion.loaders"):
        discovered = discover_documents(documents_dir)

    assert [path.relative_to(documents_dir).as_posix() for path in discovered] == [
        "nested/valid.PDF"
    ]
    assert "temporary document file" in caplog.text
    assert "unsupported file" in caplog.text


def test_password_protected_pdf_is_reported_without_attempting_passwords(
    tmp_path: Path,
) -> None:
    documents_dir = tmp_path / "documents"
    documents_dir.mkdir()
    source = documents_dir / "locked.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("test-password")
    with source.open("wb") as destination:
        writer.write(destination)

    with pytest.raises(DocumentLoadError, match="password-protected"):
        load_document(source, documents_dir)


def test_image_only_pdf_reports_ocr_limitation(tmp_path: Path) -> None:
    documents_dir = tmp_path / "documents"
    documents_dir.mkdir()
    source = documents_dir / "scan.pdf"
    make_image_only_pdf(source)

    with pytest.raises(DocumentLoadError, match="appears scanned.*OCR was not performed"):
        load_document(source, documents_dir)


def test_empty_pdf_and_empty_docx_are_actionable(tmp_path: Path) -> None:
    documents_dir = tmp_path / "documents"
    documents_dir.mkdir()
    empty_pdf = documents_dir / "empty.pdf"
    make_text_pdf(empty_pdf, ())
    empty_docx = documents_dir / "empty.docx"
    Document().save(empty_docx)

    with pytest.raises(DocumentLoadError, match="no extractable text"):
        load_document(empty_pdf, documents_dir)
    with pytest.raises(DocumentLoadError, match="no extractable paragraph or table text"):
        load_document(empty_docx, documents_dir)


def test_batch_continues_after_corrupt_file_and_persists_failure_manifest(
    tmp_path: Path,
) -> None:
    documents_dir = tmp_path / "documents"
    documents_dir.mkdir()
    valid_docx = documents_dir / "valid.docx"
    make_docx(valid_docx)
    corrupt_pdf = documents_dir / "broken.pdf"
    corrupt_pdf.write_bytes(b"not a PDF")
    manifest_path = tmp_path / "data" / "manifests" / "ingestion_manifest.json"

    batch = load_documents(documents_dir, manifest_path=manifest_path)

    assert [document.file_name for document in batch.documents] == ["valid.docx"]
    assert len(batch.failures) == 1
    assert batch.failures[0].source_path == "broken.pdf"
    assert batch.failures[0].file_hash == hashlib.sha256(corrupt_pdf.read_bytes()).hexdigest()
    assert "Could not parse PDF" in batch.failures[0].error

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    statuses = {entry["source_path"]: entry["status"] for entry in manifest["files"]}
    assert statuses == {"broken.pdf": "failed", "valid.docx": "loaded"}
    assert not list(manifest_path.parent.glob("*.tmp"))


def test_files_outside_document_root_are_rejected_before_hashing(tmp_path: Path) -> None:
    documents_dir = tmp_path / "documents"
    documents_dir.mkdir()
    outside_file = tmp_path / "outside.pdf"
    make_text_pdf(outside_file, ("Do not load this.",))

    with pytest.raises(DocumentLoadError, match="located under the documents directory"):
        load_document(outside_file, documents_dir)
